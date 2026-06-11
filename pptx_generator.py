"""
IC deck generator — produces an editable .pptx from findings data.
5 slides: Cover · Transaction Overview · Financial Scorecard · DD Findings · Recommendation
"""

import io
from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

# ── Palette ───────────────────────────────────────────────────────────────────
NAVY       = RGBColor(0x1e, 0x3a, 0x5f)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
INK        = RGBColor(0x1f, 0x29, 0x37)
GRAY       = RGBColor(0x6b, 0x72, 0x80)
LGRAY      = RGBColor(0xf3, 0xf4, 0xf6)
RULE       = RGBColor(0xe5, 0xe7, 0xeb)
GREEN      = RGBColor(0x15, 0x80, 0x3d)
AMBER      = RGBColor(0xb4, 0x53, 0x09)
RED        = RGBColor(0x99, 0x1b, 0x1b)
BLUE       = RGBColor(0x1d, 0x4e, 0xd8)
LIGHT_BLUE = RGBColor(0x93, 0xc5, 0xfd)
LIGHT_GRAY = RGBColor(0xd1, 0xd5, 0xdb)

_ASSESSMENT_COLOR = {"buy": GREEN, "conditional_buy": AMBER, "hold": AMBER, "pass": RED}
_ASSESSMENT_LABEL = {"buy": "BUY", "conditional_buy": "CONDITIONAL BUY", "hold": "HOLD", "pass": "PASS"}

# ── Dimensions ────────────────────────────────────────────────────────────────
W        = Inches(13.33)
H        = Inches(7.5)
MARGIN   = Inches(0.65)
CW       = W - 2 * MARGIN   # content width
HDR_H    = Inches(1.05)


# ── Low-level helpers ─────────────────────────────────────────────────────────

def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, l, t, w, h, fill: RGBColor, line_color: RGBColor | None = None):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = Pt(0.5)
    else:
        s.line.fill.background()
    return s


def _tf(slide, l, t, w, h) -> object:
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    return tb.text_frame


def _run(para, text, size=11, bold=False, italic=False, color=INK):
    r = para.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return r


def _p0(tf, text, size=11, bold=False, italic=False, color=INK, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size, bold, italic, color)
    return p


def _p(tf, text, size=11, bold=False, italic=False, color=INK,
        align=PP_ALIGN.LEFT, before=0):
    p = tf.add_paragraph()
    p.alignment = align
    if before:
        p.space_before = Pt(before)
    _run(p, text, size, bold, italic, color)
    return p


def _label(slide, text, top):
    tf = _tf(slide, MARGIN, top, CW, Inches(0.22))
    _p0(tf, text.upper(), size=7.5, bold=True, color=GRAY)


def _hdr(slide, title: str):
    _rect(slide, 0, 0, W, HDR_H, NAVY)
    tf = _tf(slide, MARGIN, Inches(0.28), CW, Inches(0.55))
    _p0(tf, title, size=20, bold=True, color=WHITE)
    # thin accent line at bottom of header
    _rect(slide, 0, HDR_H - Inches(0.04), W, Inches(0.04), GREEN)


# ── Slide 1: Cover ────────────────────────────────────────────────────────────

def _cover(prs, fd):
    slide = _blank(prs)
    _bg(slide, NAVY)

    dc  = fd.get("deal_context") or {}
    rec = fd.get("recommendation") or {}
    sc  = fd.get("scorecard") or {}

    assessment = rec.get("assessment", "")
    label      = _ASSESSMENT_LABEL.get(assessment, assessment.replace("_", " ").upper())
    acc_col    = _ASSESSMENT_COLOR.get(assessment, WHITE)
    composite  = sc.get("composite", 0)
    band       = sc.get("band", "")

    # Verdict pill
    _rect(slide, MARGIN, Inches(1.4), Inches(2.6), Inches(0.5), acc_col)
    tf = _tf(slide, MARGIN + Inches(0.08), Inches(1.46), Inches(2.44), Inches(0.4))
    _p0(tf, label, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Company / target name
    desc = (dc.get("company_description") or "").strip()
    name = desc.split(".")[0][:60] if desc else "Target Company"
    tf = _tf(slide, MARGIN, Inches(2.1), CW, Inches(1.1))
    _p0(tf, name, size=34, bold=True, color=WHITE)

    # Subtitle
    tf = _tf(slide, MARGIN, Inches(3.3), CW, Inches(0.35))
    _p0(tf, "INVESTMENT COMMITTEE PRESENTATION", size=10, color=LIGHT_BLUE)

    # Deal params
    parts = [
        dc.get("buyer_segment", "").replace("_", " ").title(),
        dc.get("stage", "").title(),
        dc.get("vertical", "").replace("_", " ").title(),
        f"${dc['arr_usd_m']}M ARR" if dc.get("arr_usd_m") else "",
        f"${dc['price']}M deal" if dc.get("price") else "",
    ]
    tf = _tf(slide, MARGIN, Inches(3.8), CW, Inches(0.35))
    _p0(tf, "  ·  ".join(p for p in parts if p), size=11, color=LIGHT_GRAY)

    # Score line
    tf = _tf(slide, MARGIN, Inches(4.25), CW, Inches(0.35))
    _p0(tf, f"Composite Score: {composite:.2f} / 4.00  ·  {band}", size=11, color=LIGHT_GRAY)

    # Divider
    _rect(slide, MARGIN, Inches(4.75), CW, Pt(1), GRAY)

    # Date
    tf = _tf(slide, MARGIN, Inches(6.9), CW, Inches(0.3))
    _p0(tf, date.today().strftime("%B %d, %Y"), size=9, color=GRAY, align=PP_ALIGN.RIGHT)

    # Confidentiality footer
    tf = _tf(slide, MARGIN, Inches(6.9), CW, Inches(0.3))
    _p0(tf, "CONFIDENTIAL — FOR INTERNAL USE ONLY", size=8, italic=True, color=GRAY)


# ── Slide 2: Transaction Overview ─────────────────────────────────────────────

def _transaction(prs, fd):
    slide = _blank(prs)
    _hdr(slide, "Transaction Overview")

    dc  = fd.get("deal_context") or {}
    sc  = fd.get("scorecard") or {}
    rec = fd.get("recommendation") or {}

    TOP  = HDR_H + Inches(0.3)
    HALF = (CW - Inches(0.5)) / 2
    RX   = MARGIN + HALF + Inches(0.5)

    # ── Left: deal parameters ──
    _label(slide, "Deal Parameters", TOP)
    rows = [
        ("Buyer Type",     dc.get("buyer_segment", "").replace("_", " ").title()),
        ("Target Stage",   dc.get("stage", "").title()),
        ("Vertical",       dc.get("vertical", "").replace("_", " ").title()),
        ("Size Band",      dc.get("size_band", "")),
        ("ARR",            f"${dc['arr_usd_m']}M" if dc.get("arr_usd_m") else "—"),
        ("Deal Price",     f"${dc['price']}M" if dc.get("price") else "TBD"),
        ("Return Target",  f"{dc.get('return_target','')}% IRR"),
        ("Hold Period",    f"{dc.get('hold_years','')} years"),
        ("Thesis Tags",    ", ".join(dc.get("thesis_tags") or []) or "—"),
    ]
    tf = _tf(slide, MARGIN, TOP + Inches(0.28), HALF, Inches(4.8))
    first = True
    for lbl, val in rows:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if not first:
            p.space_before = Pt(5)
        first = False
        _run(p, f"{lbl}: ", size=10, bold=True)
        _run(p, val, size=10)

    # ── Right: scorecard summary ──
    _label(slide, "Scorecard Summary", TOP)

    composite = sc.get("composite", 0)
    band      = sc.get("band", "")

    tf = _tf(slide, RX, TOP + Inches(0.28), HALF, Inches(0.75))
    _p0(tf, f"{composite:.2f} / 4.00", size=26, bold=True, color=NAVY)
    _p(tf, band, size=11, color=GRAY, before=2)

    per_family = sc.get("per_family") or {}
    _label(slide, "Score by Family", TOP + Inches(1.3))
    tf2 = _tf(slide, RX, TOP + Inches(1.55), HALF, Inches(2.0))
    first = True
    for fam, fs in per_family.items():
        if not isinstance(fs, dict):
            continue
        line = f"{fam.title()}: {fs.get('composite', 0):.2f}  ({fs.get('band', '')})"
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        if not first:
            p.space_before = Pt(4)
        first = False
        _run(p, line, size=10)

    risk_flags = fd.get("risk_flags") or []
    if risk_flags:
        _label(slide, "Risk Flags", TOP + Inches(3.05))
        tf3 = _tf(slide, RX, TOP + Inches(3.3), HALF, Inches(2.0))
        first = True
        for flag in risk_flags[:4]:
            short = flag[:80] + ("…" if len(flag) > 80 else "")
            p = tf3.paragraphs[0] if first else tf3.add_paragraph()
            if not first:
                p.space_before = Pt(4)
            first = False
            _run(p, f"• {short}", size=9, color=RED)


# ── Slide 3: Financial Scorecard ──────────────────────────────────────────────

_METRIC_ORDER = [
    ("arr_growth",         "ARR Growth",              "%"),
    ("nrr",                "Net Revenue Retention",   "%"),
    ("grr",                "Gross Revenue Retention", "%"),
    ("rule_of_40",         "Rule of 40",              ""),
    ("gross_margin",       "Gross Margin",            "%"),
    ("ebitda_margin",      "EBITDA Margin",           "%"),
    ("burn_multiple",      "Burn Multiple",           "x"),
    ("magic_number",       "Magic Number",            "x"),
    ("ltv_cac",            "LTV:CAC",                 "x"),
    ("cac_payback_months", "CAC Payback",             " mo"),
]

_BAND_COL = {
    "Strong": GREEN, "Above": BLUE, "Above Average": GREEN,
    "Below": AMBER,  "Weak": RED,   "Average": AMBER,
}


def _scorecard(prs, fd):
    slide = _blank(prs)
    _hdr(slide, "Financial Scorecard")

    gm  = fd.get("graded_metrics") or {}
    r40 = fd.get("rule_of_40")
    all_m = {**gm}
    if r40 and isinstance(r40, dict):
        all_m["rule_of_40"] = r40

    TOP   = HDR_H + Inches(0.2)
    # Column x-positions
    CX = [MARGIN, Inches(3.1), Inches(4.8), Inches(6.4), Inches(8.1), Inches(9.6), Inches(11.1)]
    HDR_LABELS = ["METRIC", "VALUE", "BAND", "TREND", "PTS", "P25 / MED / P75"]

    # Column headers
    for i, h in enumerate(HDR_LABELS):
        tf = _tf(slide, CX[i], TOP, Inches(1.8), Inches(0.25))
        _p0(tf, h, size=7.5, bold=True, color=GRAY)
    _rect(slide, MARGIN, TOP + Inches(0.26), CW, Pt(1), RULE)

    ROW_H = Inches(0.47)
    row_i = 0
    for key, name, unit in _METRIC_ORDER:
        m = all_m.get(key)
        if not m or not isinstance(m, dict):
            continue
        y = TOP + Inches(0.32) + row_i * ROW_H

        # Alternate row shading
        if row_i % 2 == 0:
            _rect(slide, MARGIN - Inches(0.05), y - Inches(0.04),
                  CW + Inches(0.1), ROW_H, LGRAY)

        band  = m.get("band", "")
        value = m.get("value", 0)
        trend = m.get("trend", "") or ""
        pts   = m.get("points", "")
        bm    = m.get("benchmark") or {}
        if not isinstance(bm, dict):
            bm = {}

        val_str = f"{value:.1f}{unit}"
        bm_str  = (f"{bm.get('q1','—')} / {bm.get('median','—')} / {bm.get('q3','—')}"
                   if bm else "—")
        t_lower = trend.lower()
        arrow   = "↑ " if "improv" in t_lower else "↓ " if "declin" in t_lower else "→ "
        trend_str = (arrow + trend) if trend else "—"
        bc = _BAND_COL.get(band, INK)

        cells = [
            (name,       INK,  True,  10),
            (val_str,    INK,  True,  10),
            (band,       bc,   True,  9),
            (trend_str,  INK,  False, 9),
            (str(pts),   bc,   True,  10),
            (bm_str,     GRAY, False, 8),
        ]
        for ci, (txt, col, bld, sz) in enumerate(cells):
            tf = _tf(slide, CX[ci], y, Inches(1.75), ROW_H - Inches(0.04))
            _p0(tf, txt, size=sz, bold=bld, color=col)
        row_i += 1


# ── Slide 4: DD Findings ──────────────────────────────────────────────────────

def _dd_findings(prs, fd):
    questions = fd.get("dd_questions") or []
    resolved  = [q for q in questions if isinstance(q, dict) and q.get("status") == "resolved"]
    if not resolved:
        return

    slide = _blank(prs)
    _hdr(slide, "Due Diligence Findings")

    allayed   = [q for q in resolved if q.get("outcome") == "allayed"]
    confirmed = [q for q in resolved if q.get("outcome") == "risk_confirmed"]
    breakers  = [q for q in resolved if q.get("outcome") == "deal_breaker"]

    TOP   = HDR_H + Inches(0.3)
    HALF  = (CW - Inches(0.4)) / 2
    RX    = MARGIN + HALF + Inches(0.4)

    def _col(x, items, accent, header_text):
        _rect(slide, x, TOP, HALF, Inches(0.38), accent)
        tf = _tf(slide, x + Inches(0.12), TOP + Inches(0.06), HALF - Inches(0.24), Inches(0.28))
        _p0(tf, header_text, size=9, bold=True, color=WHITE)

        tf2 = _tf(slide, x, TOP + Inches(0.48), HALF, Inches(5.5))
        first = True
        for q in items:
            q_text = (q.get("question") or "")[:85]
            a_text = (q.get("answer") or "")[:110]
            p = tf2.paragraphs[0] if first else tf2.add_paragraph()
            if not first:
                p.space_before = Pt(7)
            first = False
            _run(p, f"• {q_text}", size=9.5, bold=True)
            if a_text:
                p2 = tf2.add_paragraph()
                p2.space_before = Pt(1)
                _run(p2, f"   {a_text}", size=8.5, color=GRAY)

    right_label = (f"✗ Deal Breakers ({len(breakers)})  " if breakers else "") + f"⚠ Risk Confirmed ({len(confirmed)})"
    right_color = RED if breakers else AMBER

    _col(MARGIN, allayed,           GREEN,       f"✓  Allayed ({len(allayed)})")
    _col(RX,     breakers+confirmed, right_color, right_label)


# ── Slide 5: IC Recommendation ────────────────────────────────────────────────

def _recommendation(prs, fd):
    slide = _blank(prs)
    _hdr(slide, "IC Recommendation")

    rec = fd.get("recommendation") or {}
    dc  = fd.get("deal_context") or {}
    sc  = fd.get("scorecard") or {}

    assessment = rec.get("assessment", "")
    label      = _ASSESSMENT_LABEL.get(assessment, assessment.replace("_", " ").upper())
    acc_col    = _ASSESSMENT_COLOR.get(assessment, INK)
    composite  = sc.get("composite", 0)
    band       = sc.get("band", "")

    TOP = HDR_H + Inches(0.3)

    # Verdict box
    _rect(slide, MARGIN, TOP, Inches(3.2), Inches(0.65), acc_col)
    tf = _tf(slide, MARGIN + Inches(0.1), TOP + Inches(0.1), Inches(3.0), Inches(0.5))
    _p0(tf, label, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Score alongside
    tf2 = _tf(slide, MARGIN + Inches(3.4), TOP + Inches(0.08), Inches(4.5), Inches(0.55))
    _p0(tf2, f"{composite:.2f} / 4.00  ·  {band}", size=13, bold=True, color=NAVY)

    # Rationale
    rationale = (rec.get("rationale") or "")[:200]
    tf3 = _tf(slide, MARGIN, TOP + Inches(0.85), CW, Inches(0.6))
    _p0(tf3, rationale, size=10.5, color=INK)

    _rect(slide, MARGIN, TOP + Inches(1.55), CW, Pt(1), RULE)

    # Two-column: conditions + valuation
    HALF = (CW - Inches(0.5)) / 2
    RX   = MARGIN + HALF + Inches(0.5)

    conditions = rec.get("conditions") or []
    _label(slide, "Conditions", TOP + Inches(1.65))
    tf4 = _tf(slide, MARGIN, TOP + Inches(1.9), HALF, Inches(3.5))
    if conditions:
        first = True
        for i, cond in enumerate(conditions, 1):
            short = cond[:100] + ("…" if len(cond) > 100 else "")
            p = tf4.paragraphs[0] if first else tf4.add_paragraph()
            if not first:
                p.space_before = Pt(6)
            first = False
            _run(p, f"{i}. {short}", size=9.5)
    else:
        _p0(tf4, "No conditions — clean close.", size=9.5, italic=True, color=GRAY)

    _label(slide, "Valuation", TOP + Inches(1.65))
    tf5 = _tf(slide, RX, TOP + Inches(1.9), HALF, Inches(1.5))
    price = dc.get("price")
    arr   = dc.get("arr_usd_m")
    sfp   = rec.get("solve_for_price")
    if price and arr:
        multiple = round(price / arr, 1)
        _p0(tf5, f"${price}M", size=18, bold=True, color=NAVY)
        _p(tf5, f"{multiple}x ARR implied", size=10, color=GRAY, before=3)
    elif sfp and isinstance(sfp, dict):
        _p0(tf5, f"${sfp.get('low','?')}M – ${sfp.get('high','?')}M", size=14, bold=True, color=NAVY)
        _p(tf5, f"Midpoint ${sfp.get('mid','?')}M", size=10, color=GRAY, before=3)
        if sfp.get("basis"):
            _p(tf5, sfp["basis"][:80], size=9, italic=True, color=GRAY, before=3)
    else:
        _p0(tf5, "No price provided", size=10, italic=True, color=GRAY)

    # Risk flags
    risk_flags = fd.get("risk_flags") or []
    if risk_flags:
        _label(slide, "Remaining Risk Flags", TOP + Inches(3.55))
        tf6 = _tf(slide, RX, TOP + Inches(3.8), HALF, Inches(1.8))
        first = True
        for flag in risk_flags[:3]:
            short = flag[:80] + ("…" if len(flag) > 80 else "")
            p = tf6.paragraphs[0] if first else tf6.add_paragraph()
            if not first:
                p.space_before = Pt(4)
            first = False
            _run(p, f"• {short}", size=9, color=RED)


# ── Entry point ───────────────────────────────────────────────────────────────

def generate_ic_deck(findings_dict: dict) -> bytes:
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    _cover(prs, findings_dict)
    _transaction(prs, findings_dict)
    _scorecard(prs, findings_dict)
    _dd_findings(prs, findings_dict)
    _recommendation(prs, findings_dict)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
